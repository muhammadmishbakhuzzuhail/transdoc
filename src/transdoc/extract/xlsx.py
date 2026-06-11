"""Excel (.xlsx) extraction via openpyxl.

One IR block per non-empty string cell, id = ``{sheet}!{coordinate}`` (e.g. ``Sheet1!B3``).
Numbers, dates, and formulas are left alone — only text cells are translated. The renderer
reopens the workbook and writes translations back to the same coordinates, keeping styles,
column widths, merged ranges, and formulas intact.
"""

from __future__ import annotations

from ..config import Config
from ..ir import Block, BlockType, Confidence, Document
from .base import reflow_order


def extract(path: str, cfg: Config) -> Document:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=False)
    out = Document(source_path=path,
                   mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    page = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.strip() and not v.startswith("="):
                    out.blocks.append(Block(
                        id=f"{ws.title}!{cell.coordinate}", type=BlockType.PARAGRAPH,
                        page=page, text=v, confidence=Confidence(source="digital")))
        page += 1
    reflow_order(out)
    return out
