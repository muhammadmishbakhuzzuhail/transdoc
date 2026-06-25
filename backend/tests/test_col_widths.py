# SPDX-License-Identifier: AGPL-3.0-only
# Copyright (C) 2026 Muhammad Mishbakhuz Zuhail
"""Table column widths captured (docx) + rendered (docx widths + pdf colgroup)."""

from __future__ import annotations

import pytest

from transdoc.config import Config
from transdoc.ir import Cell, Table


def test_pdf_table_html_colgroup():
    from transdoc.regenerate.pdf_out import _table_html
    t = Table(rows=[[Cell(text="a"), Cell(text="b")]], col_widths=[120.0, 60.0])
    html = _table_html(t)
    assert "<colgroup>" in html and "width:120pt" in html and "width:60pt" in html


def test_table_html_accepts_bare_rows():
    from transdoc.regenerate.pdf_out import _table_html
    assert "<td" in _table_html([[Cell(text="x")]])   # back-compat: list of rows


def test_docx_extract_captures_col_widths(tmp_path):
    docx = pytest.importorskip("docx")
    from docx.shared import Pt

    from transdoc.extract.docx import extract
    from transdoc.ir import BlockType
    dd = docx.Document()
    t = dd.add_table(rows=1, cols=2)
    t.autofit = False
    t.columns[0].width = Pt(100)
    t.columns[1].width = Pt(50)
    t.cell(0, 0).text = "a"
    t.cell(0, 1).text = "b"
    p = tmp_path / "t.docx"
    dd.save(str(p))
    blk = next(b for b in extract(str(p), Config(target_lang="id")).blocks
               if b.type == BlockType.TABLE)
    assert blk.table.col_widths and round(blk.table.col_widths[0]) == 100
